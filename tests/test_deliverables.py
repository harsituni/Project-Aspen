"""The deck and the portfolio page: structure, and that their numbers match the model.

Monetary figures are real CAD.

These cannot be checked by eye in CI, so the things that actually go wrong get
asserted instead: shapes drifting off the canvas, a slide losing its speaker
notes, or a headline figure hard-coded and then drifting away from the model.
"""

from __future__ import annotations

import re
from pathlib import Path

import pytest

from src import config, deck_builder, market, montecarlo, portfolio_page, rotation, style


@pytest.fixture(scope="session")
def built(tmp_path_factory, dcf_result, mc_result, params, economics):
    """Build the deck and the page once, into a temp directory."""
    config.ensure_directories()
    rot = rotation.solve_rotations(params, economics)
    mk, _ = market.build(offline=True)
    from src import dcf as dcf_module

    if not dcf_result.grids:
        dcf_result.grids = {
            "price_vs_discount": dcf_module.sensitivity_price_vs_discount(dcf_result),
            "price_paid_vs_growth": dcf_module.sensitivity_price_paid_vs_growth(dcf_result),
        }
    figures = {p.stem: str(p) for p in config.FIGURES.glob("*.png")}
    out = tmp_path_factory.mktemp("deliverables")
    deck = deck_builder.build(rot, dcf_result, mc_result, mk, figures,
                              path=out / "deck.pptx")
    page = portfolio_page.build(rot, dcf_result, mc_result, mk, figures,
                                path=out / "page.pdf")
    return {"deck": deck, "page": page, "rot": rot, "market": mk}


# --- Deck ---------------------------------------------------------------

def _slides(path):
    from pptx import Presentation

    return Presentation(path)


def test_deck_is_six_widescreen_slides(built):
    prs = _slides(built["deck"])
    assert len(prs.slides) == 6
    ratio = prs.slide_width / prs.slide_height
    assert ratio == pytest.approx(16 / 9, rel=1e-3)


def test_every_slide_has_speaker_notes(built):
    prs = _slides(built["deck"])
    for i, slide in enumerate(prs.slides, 1):
        notes = slide.notes_slide.notes_text_frame.text.strip()
        assert len(notes) > 250, f"slide {i} has thin or missing speaker notes"


def test_no_shape_falls_off_the_canvas(built):
    prs = _slides(built["deck"])
    slack = 9144  # 0.01 inch, for rounding
    for i, slide in enumerate(prs.slides, 1):
        for shape in slide.shapes:
            if shape.left is None or shape.top is None:
                continue
            assert shape.left >= -slack, f"slide {i}: shape off the left edge"
            assert shape.top >= -slack, f"slide {i}: shape off the top edge"
            assert shape.left + shape.width <= prs.slide_width + slack, \
                f"slide {i}: shape off the right edge"
            assert shape.top + shape.height <= prs.slide_height + slack, \
                f"slide {i}: shape off the bottom edge"


def test_every_slide_is_labelled_hypothetical(built):
    """The honesty rule, enforced rather than trusted."""
    prs = _slides(built["deck"])
    for i, slide in enumerate(prs.slides, 1):
        text = " ".join(
            shape.text_frame.text for shape in slide.shapes if shape.has_text_frame
        ).lower()
        assert "hypothetical" in text, f"slide {i} is not labelled hypothetical"


def test_no_slide_mentions_a_real_asset_manager(built):
    """Nothing may imply the work came from, or represents, an employer."""
    prs = _slides(built["deck"])
    banned = ("aimco", "alberta investment management")
    for i, slide in enumerate(prs.slides, 1):
        blocks = [shape.text_frame.text for shape in slide.shapes if shape.has_text_frame]
        blocks.append(slide.notes_slide.notes_text_frame.text)
        text = " ".join(blocks).lower()
        for term in banned:
            assert term not in text, f"slide {i} names {term!r}"


def test_deck_headline_numbers_match_the_model(built, dcf_result):
    """Guards against a number being typed into the deck and then drifting."""
    prs = _slides(built["deck"])
    recommendation = prs.slides[1]
    text = " ".join(
        shape.text_frame.text for shape in recommendation.shapes if shape.has_text_frame
    )
    assert f"{dcf_result.irr:.1%}" in text
    assert style.money(dcf_result.npv) in text
    assert style.money(dcf_result.breakeven_price) in text
    assert f"{dcf_result.moic:.2f}x" in text


def test_slide_titles_are_claims_not_labels(built):
    """A takeaway title asserts something, so it should carry a number."""
    prs = _slides(built["deck"])
    for i, slide in enumerate(prs.slides, 1):
        if i == 1:
            continue  # the cover is allowed to be a plain title
        takeaway = next(
            (s.text_frame.text for s in slide.shapes
             if s.has_text_frame and s.name == "takeaway"), None
        )
        assert takeaway, f"slide {i} has no shape named 'takeaway'"
        assert re.search(r"\d", takeaway), \
            f"slide {i} takeaway has no number: {takeaway!r}"
        assert not takeaway.isupper(), f"slide {i} takeaway looks like a label"
        assert len(takeaway) > 40, f"slide {i} takeaway is too terse to be a claim"


# --- Portfolio page -------------------------------------------------------

def test_page_is_one_landscape_letter_sheet(built):
    pymupdf = pytest.importorskip("pymupdf")
    with pymupdf.open(built["page"]) as doc:
        assert len(doc) == 1
        rect = doc[0].rect
        assert rect.width / 72 == pytest.approx(11.0, abs=0.02)
        assert rect.height / 72 == pytest.approx(8.5, abs=0.02)
        assert rect.width > rect.height, "the page must be landscape"


def test_page_states_the_recommendation_and_the_caveat(built, dcf_result):
    pymupdf = pytest.importorskip("pymupdf")
    with pymupdf.open(built["page"]) as doc:
        text = doc[0].get_text().lower()
    assert "hypothetical" in text
    assert "project aspen" in text
    assert f"{dcf_result.irr:.1%}" in text
    assert style.money(dcf_result.breakeven_price).lower() in text
    assert "upper bound" in text, "the REIT-proxy caveat is missing from the page"


def test_page_contains_the_qr_code_and_four_figures(built):
    pymupdf = pytest.importorskip("pymupdf")
    with pymupdf.open(built["page"]) as doc:
        images = doc[0].get_images(full=True)
    assert len(images) >= 5, "expected four panels plus the QR code"


def test_palette_survives_greyscale():
    """Every data-encoding colour must stay separable on a mono printer."""
    clashes = style.check_print_contrast(style.SERIES_PRINT_CRITICAL)
    assert not clashes, f"colours too close in greyscale: {clashes}"


def test_github_url_is_configured():
    assert portfolio_page.GITHUB_URL.startswith("https://")
    assert "project-aspen" in portfolio_page.GITHUB_URL
